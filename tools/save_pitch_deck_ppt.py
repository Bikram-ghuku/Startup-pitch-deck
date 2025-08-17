import json
import os
import traceback
from langchain_core.tools import tool
from pptx import Presentation

@tool
def save_pitch_deck_ppt(content: str, filename: str = "pitch_deck.pptx") -> str:
    """
    Save pitch deck content to a PowerPoint file.
    
    Args:
        content: String containing the pitch deck content in structured format.
            Format should be a JSON string with slides containing:
            - title: The title of the slide
            - content: The content of the slide (can be bullet points as list or paragraphs)
            - notes: Optional presenter notes for the slide
        filename: Optional filename to save the pitch deck to (default: pitch_deck.pptx)
    
    Returns:
        Message indicating the pitch deck was saved successfully.
    """
    try:
        # Print input for debugging
        print(f"Content type: {type(content)}")

        pitch_data = json.loads(content)
            
        os.makedirs("output", exist_ok=True)
        
        if "slides" not in pitch_data:
            return "Error: Missing 'slides' key in JSON data. Please check your JSON structure."
            
        if not isinstance(pitch_data["slides"], list) or len(pitch_data["slides"]) == 0:
            return "Error: 'slides' must be a non-empty list. Please check your JSON structure."
        
        prs = Presentation()
        
        slide_count = 0
        for slide_idx, slide_data in enumerate(pitch_data.get("slides", [])):
            try:
                if "title" not in slide_data:
                    slide_data["title"] = f"Slide {slide_idx + 1}"
                
                if "content" not in slide_data:
                    slide_data["content"] = ""
                
                slide_layout = prs.slide_layouts[1]  # Title and Content layout
                slide = prs.slides.add_slide(slide_layout)
                slide_count += 1
                
                if slide.shapes.title:
                    slide.shapes.title.text = slide_data.get("title", "")
                
                if len(slide.placeholders) > 1:
                    content_shape = slide.placeholders[1]
                    content_value = slide_data.get("content", "")
                    
                    if isinstance(content_value, list):
                        text_frame = content_shape.text_frame
                        
                        for i, point in enumerate(content_value):
                            if i == 0:
                                p = text_frame.paragraphs[0]
                            else:
                                p = text_frame.add_paragraph()
                            p.text = str(point)
                            p.level = 0
                    else:
                        content_shape.text = str(content_value)
                
                notes = slide_data.get("notes")
                if notes:
                    slide_notes = slide.notes_slide
                    text_frame = slide_notes.notes_text_frame
                    text_frame.text = notes
            except Exception as slide_error:
                print(f"Error processing slide {slide_idx}: {str(slide_error)}")
                continue
        
        # Save to file
        file_path = os.path.join("output", filename)
        prs.save(file_path)
        
        return f"Pitch deck successfully saved as PowerPoint to {file_path} with {slide_count} slides"
    except Exception as e:
        # Get detailed error information
        error_details = traceback.format_exc()
        return f"Error saving PowerPoint pitch deck: {str(e)}\n\nDetails: {error_details}"